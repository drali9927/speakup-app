#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""تولید کاربرگ محتوای سطح A1 — برای تولید تصویر و صوت"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/Users/alikavyani/Desktop/speak up/content/A1-محتوا.xlsx"

STYLE = ("Children's educational illustration, soft watercolor anime style, "
         "warm friendly colors, clean simple light background, full scene, "
         "no text, no letters, no words, 4:3. Scene: ")

# ---------------------------------------------------------------- سرفصل ۳۰ درس
SYLLABUS = [
    (1, "He is a manager.", "ضمائر فاعلی و افعال to be", "سلام کن و خودت را معرفی کن"),
    (2, "It's an umbrella.", "a / an و اسم‌های جمع", "اسم چیزهای دوروبرت را بگو"),
    (3, "Who is he?", "سوالی کردن با to be، who/what", "بپرس «این کیه؟ این چیه؟»"),
    (4, "This is my sister.", "صفات ملکی و 's مالکیت", "خانواده‌ات را معرفی کن"),
    (5, "How old are you?", "اعداد و پرسش سن", "سن و تولد را بگو"),
    (6, "Where are you from?", "پرسش با where، حروف اضافه مکان", "بگو اهل کجایی"),
    (7, "She is a nurse.", "مشاغل با a/an", "درباره شغل و محل کار حرف بزن"),
    (8, "I go to school.", "حال ساده: I/you/we/they", "از کارهای روزمره‌ات بگو"),
    (9, "He works every day.", "حال ساده: he/she/it و s سوم‌شخص", "روتین دیگران را تعریف کن"),
    (10, "What time is it?", "حروف اضافه زمان: at/on/in", "ساعت و روزها را درست بگو"),
    (11, "I like music.", "like/love/hate + ing", "از علاقه‌مندی‌هایت حرف بزن"),
    (12, "Some rice, please.", "اسم شمارا و ناشمارا، some/any", "درباره غذا و نوشیدنی حرف بزن"),
    (13, "I'd like a coffee.", "would like / can I have", "مثل یک مسافر سفارش بده"),
    (14, "There is a kitchen.", "there is / there are", "خانه و اتاقت را توصیف کن"),
    (15, "The cat is under the table.", "حروف اضافه مکان", "بگو هر چیز کجاست"),
    (16, "I have got a red coat.", "have got", "لباس و رنگ‌ها را توصیف کن"),
    (17, "How much is it?", "how much / this-these", "خرید کن و قیمت بپرس"),
    (18, "It's cold today.", "it's + صفت، آب‌وهوا", "درباره هوا و فصل‌ها حرف بزن"),
    (19, "I can swim.", "can / can't", "بگو چه کارهایی بلدی"),
    (20, "I have a headache.", "have got برای بیماری، should", "حال جسمت را توضیح بده"),
    (21, "Turn left, please.", "جملات امری", "مسیر بپرس و راهنمایی کن"),
    (22, "She is reading now.", "حال استمراری", "بگو همین حالا چه خبر است"),
    (23, "I usually walk.", "حال ساده در برابر استمراری", "فرق عادت و «همین الان» را بفهم"),
    (24, "I was at home.", "گذشته to be: was/were", "بگو دیروز کجا بودی"),
    (25, "We played football.", "گذشته ساده: افعال باقاعده", "از دیروزت تعریف کن"),
    (26, "I went to Tehran.", "گذشته ساده: افعال بی‌قاعده", "یک سفر گذشته را تعریف کن"),
    (27, "Did you like it?", "سوالی گذشته با did", "از سفر و تعطیلات سؤال بپرس"),
    (28, "I'm going to study.", "be going to برای آینده", "از برنامه‌های آینده‌ات بگو"),
    (29, "Tehran is bigger.", "صفات تفضیلی", "دو چیز را با هم مقایسه کن"),
    (30, "The best day.", "صفات عالی و مرور کلی", "بهترین‌ها را بگو و A1 را جمع کن"),
]

# ---------------------------------------------------------------- واژگان
# (واژه, بخش کلام, IPA, فارسی, جمله انگلیسی, ترجمه, صحنه تصویر)
V = {}

V[1] = [
 ("hello","interjection","/həˈləʊ/","سلام","Hello! My name is Sara.","سلام! اسم من ساراست.","two colleagues in their early thirties greeting each other with a friendly wave in a modern office lobby"),
 ("hi","interjection","/haɪ/","سلام (خودمانی)","Hi! How are you?","سلام! حالت چطوره؟","two friends around thirty greeting each other casually across a cafe table"),
 ("goodbye","interjection","/ˌɡʊdˈbaɪ/","خداحافظ","Goodbye! See you tomorrow.","خداحافظ! فردا می‌بینمت.","a woman in her early thirties waving goodbye to a friend from a train platform"),
 ("morning","noun","/ˈmɔːnɪŋ/","صبح","Good morning, teacher!","صبح بخیر، معلم!","a man in his early thirties holding a coffee mug by a sunlit window in the early morning"),
 ("afternoon","noun","/ˌɑːftəˈnuːn/","بعدازظهر","Good afternoon! Come in, please.","بعدازظهر بخیر! بفرمایید داخل.","a woman in her thirties welcoming a guest into a bright apartment in warm afternoon light"),
 ("evening","noun","/ˈiːvnɪŋ/","عصر","Good evening, everyone.","عصر بخیر، به همگی.","three adults in their thirties greeting each other around a dinner table in warm evening light"),
 ("night","noun","/naɪt/","شب","Good night! Sleep well.","شب بخیر! خوب بخوابی.","a woman in her early thirties switching off a bedside lamp at night, city lights through the window"),
 ("name","noun","/neɪm/","نام","My name is Ali.","اسم من علی است.","a man in his thirties introducing himself with a hand on his chest at a professional networking event"),
 ("teacher","noun","/ˈtiːtʃə/","معلم","She is my English teacher.","او معلم انگلیسی من است.","a confident woman in her thirties teaching a small group of adults beside a whiteboard"),
 ("student","noun","/ˈstjuːdnt/","زبان‌آموز","I am a student.","من زبان‌آموز هستم.","a man around thirty studying with a laptop and notebook at a library table"),
 ("manager","noun","/ˈmænɪdʒə/","مدیر","He is a manager.","او یک مدیر است.","a man in his thirties in smart-casual clothes leading a small meeting in a modern office"),
 ("doctor","noun","/ˈdɒktə/","پزشک","My father is a doctor.","پدر من پزشک است.","a doctor in their thirties in a white coat with a stethoscope in a bright clinic"),
 ("friend","noun","/frend/","دوست","This is my friend, Reza.","این دوست من، رضاست.","two friends around thirty walking and laughing together on a city street"),
 ("please","adverb","/pliːz/","لطفاً","Please sit down.","لطفاً بنشینید.","a man in his thirties politely gesturing toward an empty chair, offering a seat"),
 ("thanks","noun","/θæŋks/","ممنون","Thanks for your help!","ممنون بابت کمکت!","a woman in her thirties receiving a folder from a colleague with a grateful smile"),
 ("yes","adverb","/jes/","بله","Yes, I am a teacher.","بله، من معلم هستم.","a man in his thirties nodding in agreement, a large green check mark floating beside him"),
 ("no","adverb","/nəʊ/","نه","No, he is not a doctor.","نه، او پزشک نیست.","a woman in her thirties politely shaking her head, a large red cross mark floating beside her"),
 ("sorry","adjective","/ˈsɒri/","ببخشید","Sorry, I am late.","ببخشید، دیر کردم.","a man in his early thirties arriving late to a meeting room with an apologetic expression"),
 ("welcome","adjective","/ˈwelkəm/","خوش‌آمد","Welcome to our class!","به کلاس ما خوش آمدید!","a woman in her thirties holding a door open and welcoming a guest into an office"),
 ("meet","verb","/miːt/","ملاقات کردن","Nice to meet you.","از ملاقات شما خوشبختم.","two professionals around thirty shaking hands and smiling at a business meeting"),
]

V[2] = [
 ("umbrella","noun","/ʌmˈbrelə/","چتر","It's an umbrella.","این یک چتر است.","a colorful open umbrella standing alone"),
 ("book","noun","/bʊk/","کتاب","This is a book.","این یک کتاب است.","a single open book on a wooden table"),
 ("pen","noun","/pen/","خودکار","I have a blue pen.","من یک خودکار آبی دارم.","a blue ballpoint pen on a white notebook"),
 ("pencil","noun","/ˈpensl/","مداد","The pencil is on the desk.","مداد روی میز است.","a sharpened pencil lying on a wooden desk next to a notebook"),
 ("bag","noun","/bæɡ/","کیف","My bag is heavy.","کیف من سنگین است.","a leather work bag with a laptop and notebook beside it on a chair"),
 ("table","noun","/ˈteɪbl/","میز","The book is on the table.","کتاب روی میز است.","a simple wooden table in a bright room"),
 ("chair","noun","/tʃeə/","صندلی","This chair is new.","این صندلی نو است.","a single wooden chair beside a window"),
 ("door","noun","/dɔː/","در","Please close the door.","لطفاً در را ببندید.","a wooden door slightly open in a hallway"),
 ("window","noun","/ˈwɪndəʊ/","پنجره","The window is open.","پنجره باز است.","an open window with sunlight and curtains"),
 ("key","noun","/kiː/","کلید","Where is my key?","کلید من کجاست؟","a small metal key on a table"),
 ("phone","noun","/fəʊn/","تلفن","My phone is new.","تلفن من نو است.","a modern smartphone lying on a desk"),
 ("watch","noun","/wɒtʃ/","ساعت مچی","This is a nice watch.","این ساعت قشنگی است.","a wristwatch on a person's wrist"),
 ("box","noun","/bɒks/","جعبه","What is in the box?","داخل جعبه چیست؟","a closed cardboard box on the floor"),
 ("cup","noun","/kʌp/","فنجان","I want a cup of tea.","یک فنجان چای می‌خواهم.","a warm cup of tea with steam rising"),
 ("ball","noun","/bɔːl/","توپ","The ball is red.","توپ قرمز است.","a red ball on green grass"),
 ("car","noun","/kɑː/","خودرو","That is my car.","آن خودروی من است.","a small family car parked on a street"),
 ("house","noun","/haʊs/","خانه","Our house is small.","خانه ما کوچک است.","a small friendly house with a garden"),
 ("tree","noun","/triː/","درخت","There is a tree in the garden.","یک درخت در باغ است.","a large green tree in a sunny garden"),
 ("apple","noun","/ˈæpl/","سیب","I eat an apple every day.","هر روز یک سیب می‌خورم.","a shiny red apple on a table"),
 ("computer","noun","/kəmˈpjuːtə/","رایانه","She has a new computer.","او یک رایانه جدید دارد.","a laptop computer open on a desk"),
]

V[3] = [
 ("who","pronoun","/huː/","چه کسی","Who is he?","او کیست؟","a person with a question mark above their head looking at a stranger"),
 ("what","pronoun","/wɒt/","چه","What is your name?","اسم شما چیست؟","two people talking, one with a question mark above their head"),
 ("where","adverb","/weə/","کجا","Where is the book?","کتاب کجاست؟","an adult looking around a room searching for something, a question mark above their head"),
 ("man","noun","/mæn/","مرد","That man is my uncle.","آن مرد عموی من است.","a friendly adult man standing and smiling"),
 ("woman","noun","/ˈwʊmən/","زن","The woman is a doctor.","آن زن پزشک است.","a friendly adult woman standing and smiling"),
 ("boy","noun","/bɔɪ/","پسر","The boy is my brother.","آن پسر برادر من است.","a young boy standing with a backpack"),
 ("girl","noun","/ɡɜːl/","دختر","This girl is my friend.","این دختر دوست من است.","a young girl standing and smiling"),
 ("child","noun","/tʃaɪld/","کودک","The child is happy.","کودک خوشحال است.","a small happy child playing"),
 ("people","noun","/ˈpiːpl/","مردم","Many people are here.","افراد زیادی اینجا هستند.","a small group of diverse people standing together"),
 ("tall","adjective","/tɔːl/","قدبلند","He is very tall.","او خیلی قدبلند است.","a tall adult standing beside a shorter adult, clear height contrast"),
 ("short","adjective","/ʃɔːt/","کوتاه‌قد","She is short.","او کوتاه‌قد است.","a short adult standing beside a taller adult, clear height contrast"),
 ("young","adjective","/jʌŋ/","جوان","My sister is young.","خواهرم جوان است.","a person in their late twenties standing beside a person in their seventies"),
 ("old","adjective","/əʊld/","پیر","My grandfather is old.","پدربزرگم پیر است.","a kind man in his seventies with grey hair, smiling"),
 ("happy","adjective","/ˈhæpi/","خوشحال","I am happy today.","امروز خوشحالم.","an adult in their early thirties with a genuine joyful smile"),
 ("sad","adjective","/sæd/","ناراحت","Why are you sad?","چرا ناراحتی؟","an adult in their thirties sitting alone by a window, looking down"),
 ("new","adjective","/njuː/","نو","This is a new bag.","این یک کیف نو است.","a brand new leather bag with the price tag still attached"),
 ("nice","adjective","/naɪs/","خوب","She is a nice teacher.","او معلم خوبی است.","a friendly adult in their thirties offering a helping hand"),
 ("this","determiner","/ðɪs/","این","This is my book.","این کتاب من است.","an adult hand pointing at an object close to the viewer"),
 ("that","determiner","/ðæt/","آن","That is your car.","آن خودروی توست.","an adult pointing at a distant building across a street"),
 ("here","adverb","/hɪə/","اینجا","Come here, please.","لطفاً بیا اینجا.","an adult gesturing come-here toward someone across a room"),
]

V[4] = [
 ("family","noun","/ˈfæməli/","خانواده","I love my family.","خانواده‌ام را دوست دارم.","three generations of a family standing together, adults and one grandparent"),
 ("mother","noun","/ˈmʌðə/","مادر","My mother is a teacher.","مادرم معلم است.","a woman in her late fifties smiling warmly beside her adult daughter"),
 ("father","noun","/ˈfɑːðə/","پدر","His father is a driver.","پدر او راننده است.","a man in his late fifties standing proudly beside his adult son"),
 ("parent","noun","/ˈpeərənt/","والد","My parents are at home.","والدینم خانه هستند.","an older couple standing together with their adult child"),
 ("sister","noun","/ˈsɪstə/","خواهر","This is my sister.","این خواهر من است.","two adult sisters in their thirties sitting together on a sofa"),
 ("brother","noun","/ˈbrʌðə/","برادر","My brother is ten.","برادرم ده ساله است.","two adult brothers in their thirties talking together outdoors"),
 ("son","noun","/sʌn/","پسر (فرزند)","Their son is a student.","پسرشان دانش‌آموز است.","a father in his sixties with his adult son beside him"),
 ("daughter","noun","/ˈdɔːtə/","دختر (فرزند)","She is my daughter.","او دختر من است.","a mother in her sixties with her adult daughter beside her"),
 ("wife","noun","/waɪf/","همسر (زن)","His wife is a nurse.","همسر او پرستار است.","a married couple standing side by side"),
 ("husband","noun","/ˈhʌzbənd/","همسر (مرد)","Her husband is a cook.","همسر او آشپز است.","a married couple smiling together"),
 ("grandmother","noun","/ˈɡrænmʌðə/","مادربزرگ","My grandmother makes good food.","مادربزرگم غذای خوبی می‌پزد.","a kind grandmother with white hair offering a plate of food"),
 ("grandfather","noun","/ˈɡrænfɑːðə/","پدربزرگ","My grandfather is seventy.","پدربزرگم هفتاد ساله است.","a kind grandfather with a walking stick, smiling"),
 ("aunt","noun","/ɑːnt/","عمه/خاله","My aunt lives in Shiraz.","خاله‌ام در شیراز زندگی می‌کند.","a woman in her fifties warmly greeting her adult niece"),
 ("uncle","noun","/ˈʌŋkl/","عمو/دایی","Her uncle is a farmer.","عموی او کشاورز است.","a man in his fifties warmly greeting his adult nephew"),
 ("cousin","noun","/ˈkʌzn/","پسرعمو/دخترخاله","My cousin is my friend too.","پسرعمویم دوستم هم هست.","two cousins in their thirties talking and laughing together"),
 ("baby","noun","/ˈbeɪbi/","نوزاد","The baby is sleeping.","نوزاد خوابیده است.","a peaceful sleeping baby in a crib"),
 ("married","adjective","/ˈmærid/","متأهل","My sister is married.","خواهرم متأهل است.","a couple with wedding rings holding hands"),
 ("single","adjective","/ˈsɪŋɡl/","مجرد","He is single.","او مجرد است.","one person standing alone, content and smiling"),
 ("my","determiner","/maɪ/","مالِ من","This is my house.","این خانه من است.","an adult standing proudly in front of their own apartment door"),
 ("our","determiner","/ˈaʊə/","مالِ ما","Our school is big.","مدرسه ما بزرگ است.","a group of adult colleagues standing in front of their office building"),
]

V[5] = [
 ("number","noun","/ˈnʌmbə/","عدد","What is your phone number?","شماره تلفنت چند است؟","an adult counting small wooden blocks laid out on a table"),
 ("age","noun","/eɪdʒ/","سن","What is your age?","سن شما چند است؟","two adults of clearly different ages standing side by side"),
 ("year","noun","/jɪə/","سال","I am twenty years old.","من بیست ساله هستم.","a wall calendar with twelve blank month pages, no writing"),
 ("birthday","noun","/ˈbɜːθdeɪ/","تولد","Today is my birthday!","امروز تولد من است!","an adult in their thirties blowing out candles on a birthday cake"),
 ("first","ordinal","/fɜːst/","اول","This is my first day.","این اولین روز من است.","a runner crossing a finish line ahead of the others"),
 ("second","ordinal","/ˈsekənd/","دوم","She is in the second class.","او در کلاس دوم است.","an adult standing on the second step of a winners podium"),
 ("third","ordinal","/θɜːd/","سوم","He lives on the third floor.","او در طبقه سوم زندگی می‌کند.","an adult standing on the third step of a winners podium"),
 ("many","determiner","/ˈmeni/","زیاد","Many students are here.","دانش‌آموزان زیادی اینجا هستند.","a large crowd of students in a schoolyard"),
 ("few","determiner","/fjuː/","کم","Only a few people came.","فقط چند نفر آمدند.","three people standing in a large empty room"),
 ("hundred","number","/ˈhʌndrəd/","صد","There are a hundred pages.","صد صفحه دارد.","a thick book with many pages fanned out"),
 ("thousand","number","/ˈθaʊznd/","هزار","This city has a thousand shops.","این شهر هزار مغازه دارد.","a wide city view with countless small buildings"),
 ("count","verb","/kaʊnt/","شمردن","Can you count to ten?","می‌توانی تا ده بشماری؟","an adult counting on their fingers, thoughtful expression"),
 ("address","noun","/əˈdres/","نشانی","What is your address?","نشانی شما چیست؟","an adult holding a plain envelope in front of an apartment door"),
 ("class","noun","/klɑːs/","کلاس","Our class has thirty students.","کلاس ما سی دانش‌آموز دارد.","an adult language class, students at desks facing a whiteboard"),
 ("group","noun","/ɡruːp/","گروه","We work in a group.","ما گروهی کار می‌کنیم.","four adults working together around one table"),
 ("party","noun","/ˈpɑːti/","مهمانی","There is a party tonight.","امشب یک مهمانی است.","a cheerful adult birthday gathering with balloons and a cake"),
 ("gift","noun","/ɡɪft/","هدیه","This gift is for you.","این هدیه برای توست.","a wrapped gift box with a ribbon"),
 ("candle","noun","/ˈkændl/","شمع","There are ten candles.","ده شمع هست.","ten lit candles on a birthday cake"),
 ("million","number","/ˈmɪljən/","میلیون","The city has two million people.","شهر دو میلیون نفر جمعیت دارد.","a wide aerial view of a large city at dusk"),
 ("date","noun","/deɪt/","تاریخ","What is the date today?","امروز چندم است؟","a wall calendar with one day circled"),
]

V[6] = [
 ("country","noun","/ˈkʌntri/","کشور","Iran is a big country.","ایران کشور بزرگی است.","a world map with one country highlighted"),
 ("city","noun","/ˈsɪti/","شهر","Tehran is a large city.","تهران شهر بزرگی است.","a city skyline with tall buildings"),
 ("world","noun","/wɜːld/","جهان","The world is beautiful.","جهان زیباست.","a globe of the earth surrounded by children"),
 ("map","noun","/mæp/","نقشه","Look at the map.","به نقشه نگاه کن.","an open paper map on a table"),
 ("language","noun","/ˈlæŋɡwɪdʒ/","زبان","English is a world language.","انگلیسی یک زبان جهانی است.","speech bubbles with different scripts around a globe"),
 ("from","preposition","/frɒm/","اهلِ","I am from Iran.","من اهل ایران هستم.","a person standing on a map pointing at their country"),
 ("live","verb","/lɪv/","زندگی کردن","We live in Tehran.","ما در تهران زندگی می‌کنیم.","a family standing in front of their home in a city"),
 ("near","preposition","/nɪə/","نزدیک","My school is near my house.","مدرسه‌ام نزدیک خانه‌ام است.","two buildings very close to each other"),
 ("far","adjective","/fɑː/","دور","The airport is far.","فرودگاه دور است.","a long road stretching to a distant building"),
 ("big","adjective","/bɪɡ/","بزرگ","This is a big city.","این شهر بزرگی است.","a large elephant beside a small mouse for contrast"),
 ("small","adjective","/smɔːl/","کوچک","We live in a small town.","ما در شهر کوچکی زندگی می‌کنیم.","a tiny village with a few houses"),
 ("beautiful","adjective","/ˈbjuːtɪfl/","زیبا","Shiraz is beautiful.","شیراز زیباست.","a beautiful garden with flowers and a fountain"),
 ("capital","noun","/ˈkæpɪtl/","پایتخت","Tehran is the capital.","تهران پایتخت است.","a star marking a capital city on a map"),
 ("flag","noun","/flæɡ/","پرچم","Every country has a flag.","هر کشور یک پرچم دارد.","a generic colorful flag waving on a pole"),
 ("travel","verb","/ˈtrævl/","سفر کردن","I like to travel.","دوست دارم سفر کنم.","a person with a suitcase walking toward an airplane"),
 ("visit","verb","/ˈvɪzɪt/","دیدن کردن","We visit our grandmother.","به دیدن مادربزرگمان می‌رویم.","a family arriving at a grandmother's door with flowers"),
 ("foreign","adjective","/ˈfɒrən/","خارجی","She speaks a foreign language.","او یک زبان خارجی صحبت می‌کند.","a person reading a book with unfamiliar script"),
 ("north","noun","/nɔːθ/","شمال","We live in the north.","ما در شمال زندگی می‌کنیم.","a compass with the north direction highlighted"),
 ("south","noun","/saʊθ/","جنوب","The sea is in the south.","دریا در جنوب است.","a compass with the south direction highlighted"),
 ("village","noun","/ˈvɪlɪdʒ/","روستا","My uncle lives in a village.","عمویم در روستا زندگی می‌کند.","a peaceful small village with green fields"),
]

V[7] = [
 ("job","noun","/dʒɒb/","شغل","What is your job?","شغل شما چیست؟","several people in different work uniforms standing together"),
 ("work","noun","/wɜːk/","کار","I go to work at eight.","ساعت هشت سر کار می‌روم.","a person walking into an office building in the morning"),
 ("office","noun","/ˈɒfɪs/","دفتر","Her office is on floor two.","دفتر او در طبقه دوم است.","a bright modern office with desks and computers"),
 ("nurse","noun","/nɜːs/","پرستار","She is a nurse.","او پرستار است.","a friendly nurse in a hospital corridor"),
 ("engineer","noun","/ˌendʒɪˈnɪə/","مهندس","My brother is an engineer.","برادرم مهندس است.","an engineer with a helmet and blueprints at a site"),
 ("driver","noun","/ˈdraɪvə/","راننده","He is a taxi driver.","او راننده تاکسی است.","a smiling driver sitting in a taxi"),
 ("cook","noun","/kʊk/","آشپز","The cook makes good food.","آشپز غذای خوبی می‌پزد.","a chef in a white hat cooking in a kitchen"),
 ("farmer","noun","/ˈfɑːmə/","کشاورز","The farmer works outside.","کشاورز بیرون کار می‌کند.","a farmer working in a green field with a hat"),
 ("shop","noun","/ʃɒp/","مغازه","There is a shop near us.","نزدیک ما یک مغازه هست.","a small friendly corner shop with a window display"),
 ("factory","noun","/ˈfæktri/","کارخانه","He works in a factory.","او در کارخانه کار می‌کند.","a factory building with smoke stacks"),
 ("hospital","noun","/ˈhɒspɪtl/","بیمارستان","The hospital is big.","بیمارستان بزرگ است.","a hospital building with an ambulance in front"),
 ("school","noun","/skuːl/","مدرسه","Our school is new.","مدرسه ما نو است.","a modern school building seen from the street on a sunny day"),
 ("company","noun","/ˈkʌmpəni/","شرکت","She works for a company.","او در یک شرکت کار می‌کند.","a group of colleagues in a meeting room"),
 ("boss","noun","/bɒs/","رئیس","My boss is very kind.","رئیس من خیلی مهربان است.","a friendly manager talking with an employee"),
 ("money","noun","/ˈmʌni/","پول","I need some money.","کمی پول لازم دارم.","banknotes and coins on a table"),
 ("busy","adjective","/ˈbɪzi/","مشغول","I am busy today.","امروز مشغولم.","a person surrounded by many tasks and papers"),
 ("hard","adjective","/hɑːd/","سخت","This job is hard.","این کار سخت است.","a person lifting a heavy box with effort"),
 ("easy","adjective","/ˈiːzi/","آسان","The test was easy.","آزمون آسان بود.","a person relaxing while easily solving a puzzle"),
 ("worker","noun","/ˈwɜːkə/","کارگر","He is a worker in a factory.","او کارگر کارخانه است.","a worker in their thirties in a clean workshop"),
 ("police officer","noun","/pəˈliːs ˈɒfɪsə/","افسر پلیس","The police officer helps people.","افسر پلیس به مردم کمک می‌کند.","a friendly police officer helping a child cross a street"),
]

V[8] = [
 ("wake","verb","/weɪk/","بیدار شدن","I wake up at six.","ساعت شش بیدار می‌شوم.","an adult in their thirties stretching and waking up in bed at sunrise"),
 ("sleep","verb","/sliːp/","خوابیدن","I sleep at ten.","ساعت ده می‌خوابم.","an adult in their thirties sleeping peacefully in bed at night"),
 ("eat","verb","/iːt/","خوردن","We eat breakfast together.","با هم صبحانه می‌خوریم.","two adults eating breakfast together at a kitchen table"),
 ("drink","verb","/drɪŋk/","نوشیدن","I drink tea every morning.","هر صبح چای می‌نوشم.","an adult in their thirties drinking tea from a glass by a window"),
 ("go","verb","/ɡəʊ/","رفتن","I go to school.","به مدرسه می‌روم.","an adult walking toward an office building with a work bag"),
 ("come","verb","/kʌm/","آمدن","Come to my house.","به خانه من بیا.","an adult arriving at an open front door, welcomed inside"),
 ("study","verb","/ˈstʌdi/","درس خواندن","I study English.","انگلیسی می‌خوانم.","an adult studying with a book and laptop under a desk lamp at night"),
 ("play","verb","/pleɪ/","بازی کردن","The children play football.","بچه‌ها فوتبال بازی می‌کنند.","adults in their thirties playing football together in a park"),
 ("read","verb","/riːd/","خواندن","I read a book at night.","شب کتاب می‌خوانم.","an adult reading a book in a comfortable armchair"),
 ("write","verb","/raɪt/","نوشتن","She writes a letter.","او نامه می‌نویسد.","an adult writing in a notebook with a pen at a desk"),
 ("ask","verb","/ɑːsk/","پرسیدن","I ask a question.","سوالی می‌پرسم.","a man in his early thirties raising his hand to ask in a small group"),
 ("listen","verb","/ˈlɪsn/","گوش دادن","I listen to music.","به موسیقی گوش می‌دهم.","an adult wearing headphones, enjoying music with eyes closed"),
 ("walk","verb","/wɔːk/","قدم زدن","We walk to the park.","تا پارک پیاده می‌رویم.","two adults walking together along a tree-lined city path"),
 ("run","verb","/rʌn/","دویدن","He runs every morning.","او هر صبح می‌دود.","an adult jogging in a park at sunrise"),
 ("sit","verb","/sɪt/","نشستن","Please sit here.","لطفاً اینجا بنشینید.","an adult sitting down on a chair at a cafe table"),
 ("stand","verb","/stænd/","ایستادن","Stand up, please.","لطفاً بایستید.","an adult standing up from a chair in a meeting room"),
 ("open","verb","/ˈəʊpən/","باز کردن","Open your books.","کتاب‌هایتان را باز کنید.","adult hands opening a book on a table"),
 ("close","verb","/kləʊz/","بستن","Close the window, please.","لطفاً پنجره را ببندید.","adult hands closing a window"),
 ("help","verb","/help/","کمک کردن","Can you help me?","می‌توانی کمکم کنی؟","one adult helping another carry a heavy box"),
 ("call","verb","/kɔːl/","زنگ زدن","I call my mother every day.","هر روز به مادرم زنگ می‌زنم.","a woman in her early thirties speaking warmly on a phone by a window"),
]

V[9] = [
 ("always","adverb","/ˈɔːlweɪz/","همیشه","He always drinks tea.","او همیشه چای می‌نوشد.","an adult doing the same daily action, seven identical small panels"),
 ("usually","adverb","/ˈjuːʒuəli/","معمولاً","She usually walks to work.","او معمولاً پیاده سر کار می‌رود.","an adult doing a daily action, five of seven small panels filled"),
 ("often","adverb","/ˈɒfn/","اغلب","We often play football.","ما اغلب فوتبال بازی می‌کنیم.","an adult doing a daily action, four of seven small panels filled"),
 ("sometimes","adverb","/ˈsʌmtaɪmz/","گاهی","I sometimes read at night.","گاهی شب‌ها مطالعه می‌کنم.","an adult doing a daily action, two of seven small panels filled"),
 ("never","adverb","/ˈnevə/","هرگز","He never eats meat.","او هرگز گوشت نمی‌خورد.","an adult declining an action, a large red cross mark beside them"),
 ("every day","phrase","/ˈevri deɪ/","هر روز","She studies every day.","او هر روز درس می‌خواند.","one adult repeating the same morning routine across several small panels"),
 ("week","noun","/wiːk/","هفته","I work five days a week.","هفته‌ای پنج روز کار می‌کنم.","a blank weekly planner page with seven empty columns, no writing"),
 ("weekend","noun","/ˌwiːkˈend/","آخر هفته","We rest at the weekend.","آخر هفته استراحت می‌کنیم.","two adults relaxing at home on a weekend morning"),
 ("home","noun","/həʊm/","خانه","I am at home now.","الان خانه هستم.","a warm apartment interior with soft evening light"),
 ("breakfast","noun","/ˈbrekfəst/","صبحانه","Breakfast is at seven.","صبحانه ساعت هفت است.","a breakfast table with bread, cheese, herbs and tea"),
 ("lunch","noun","/lʌntʃ/","ناهار","We have lunch at one.","ساعت یک ناهار می‌خوریم.","a lunch plate with rice and vegetables on a table"),
 ("dinner","noun","/ˈdɪnə/","شام","Dinner is ready.","شام آماده است.","two adults having dinner together in warm evening light"),
 ("early","adjective","/ˈɜːli/","زود","He gets up early.","او زود بیدار می‌شود.","an adult already awake and dressed at sunrise"),
 ("late","adjective","/leɪt/","دیر","Don't be late!","دیر نکن!","an adult hurrying while glancing at a wristwatch"),
 ("together","adverb","/təˈɡeðə/","با هم","We study together.","با هم درس می‌خوانیم.","two adults studying side by side at a shared table"),
 ("alone","adverb","/əˈləʊn/","تنها","She lives alone.","او تنها زندگی می‌کند.","one adult sitting alone in a quiet room with a book"),
 ("again","adverb","/əˈɡen/","دوباره","Say it again, please.","لطفاً دوباره بگو.","a circular arrow around an adult repeating an action"),
 ("ready","adjective","/ˈredi/","آماده","Are you ready?","آماده‌ای؟","an adult with a work bag on their shoulder, ready at the front door"),
 ("finish","verb","/ˈfɪnɪʃ/","تمام کردن","I finish work at five.","ساعت پنج کارم تمام می‌شود.","an adult closing a laptop with a satisfied expression"),
 ("start","verb","/stɑːt/","شروع کردن","Class starts at eight.","کلاس ساعت هشت شروع می‌شود.","an adult at a starting line, ready to begin a run"),
]

V[10] = [
 ("time","noun","/taɪm/","زمان","What time is it?","ساعت چند است؟","an adult glancing at a large wall clock, thoughtful expression"),
 ("hour","noun","/ˈaʊə/","ساعت (مدت)","I study for one hour.","یک ساعت درس می‌خوانم.","an hourglass with sand flowing on a wooden desk"),
 ("minute","noun","/ˈmɪnɪt/","دقیقه","Wait five minutes.","پنج دقیقه صبر کن.","a stopwatch held in an adult hand"),
 ("clock","noun","/klɒk/","ساعت دیواری","The clock is on the wall.","ساعت روی دیوار است.","a round minimalist wall clock on a plain wall"),
 ("day","noun","/deɪ/","روز","Today is a good day.","امروز روز خوبی است.","a bright sunny day over a quiet city street"),
 ("month","noun","/mʌnθ/","ماه","There are twelve months.","دوازده ماه وجود دارد.","a blank wall calendar with twelve empty month pages, no writing"),
 ("Monday","noun","/ˈmʌndeɪ/","دوشنبه","I work on Monday.","دوشنبه کار می‌کنم.","an adult starting the work week at a desk on a bright morning"),
 ("Tuesday","noun","/ˈtjuːzdeɪ/","سه‌شنبه","We meet on Tuesday.","سه‌شنبه همدیگر را می‌بینیم.","two adults in a small morning meeting at an office table"),
 ("Wednesday","noun","/ˈwenzdeɪ/","چهارشنبه","The class is on Wednesday.","کلاس چهارشنبه است.","an adult attending an evening language class midweek"),
 ("Thursday","noun","/ˈθɜːzdeɪ/","پنج‌شنبه","He comes on Thursday.","او پنج‌شنبه می‌آید.","an adult finishing work and packing a bag at the end of the day"),
 ("Friday","noun","/ˈfraɪdeɪ/","جمعه","Friday is a rest day.","جمعه روز استراحت است.","an adult relaxing at home with tea on a restful morning"),
 ("Saturday","noun","/ˈsætədeɪ/","شنبه","School starts on Saturday.","مدرسه شنبه شروع می‌شود.","an adult heading out to work early on a bright morning"),
 ("Sunday","noun","/ˈsʌndeɪ/","یکشنبه","I visit my aunt on Sunday.","یکشنبه به دیدن خاله‌ام می‌روم.","two adults visiting family, arriving at a door with flowers"),
 ("today","adverb","/təˈdeɪ/","امروز","Today is Monday.","امروز دوشنبه است.","an adult marking the current day on a blank wall planner"),
 ("tomorrow","adverb","/təˈmɒrəʊ/","فردا","See you tomorrow!","فردا می‌بینمت!","an adult looking ahead at a blank planner, one arrow pointing forward"),
 ("yesterday","adverb","/ˈjestədeɪ/","دیروز","Yesterday was hot.","دیروز گرم بود.","an adult looking back over a shoulder, one arrow pointing backward"),
 ("now","adverb","/naʊ/","الان","I am busy now.","الان مشغولم.","an adult pointing at a wall clock in the present moment"),
 ("o'clock","adverb","/əˈklɒk/","ساعتِ","It is three o'clock.","ساعت سه است.","a wall clock showing an exact hour, hands straight up"),
 ("half","noun","/hɑːf/","نیم","It's half past two.","دو و نیم است.","a wall clock with the minute hand pointing straight down"),
 ("last","adjective","/lɑːst/","گذشته، قبلی","I saw him last week.","هفته گذشته دیدمش.","a wall calendar with the previous week marked"),
]

# ---------------------------------------------------------------- ساخت فایل
HDR  = PatternFill("solid", fgColor="2F3E7E")
HDRF = Font(color="FFFFFF", bold=True, size=11)
ALT  = PatternFill("solid", fgColor="EEF1FA")
TODO = PatternFill("solid", fgColor="FFF4CC")
THIN = Side(style="thin", color="C8CEE0")
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def sheet(wb, title, headers, widths, rtl=True):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = rtl
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill, c.font, c.border = HDR, HDRF, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 34
    return ws

def finish(ws, ncols, wrap_cols=()):
    for r in range(2, ws.max_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.border = BOX
            cell.alignment = Alignment(vertical="center", wrap_text=(c in wrap_cols))
            if r % 2 == 0:
                cell.fill = ALT

wb = Workbook(); wb.remove(wb.active)

# --- راهنما
g = wb.create_sheet("راهنما"); g.sheet_view.rightToLeft = True
g.column_dimensions["A"].width = 22; g.column_dimensions["B"].width = 110
rows = [
 ("کاربرگ محتوای سطح A1", ""),
 ("", ""),
 ("این فایل چیست؟", "ستون فقرات محتوای سطح A1. هر ردیف یک قلم محتواست که باید ساخته شود."),
 ("کار شما", "فقط ستون «پرامپت تصویر» را کپی کنید، به هوش مصنوعی بدهید، و تصویر خروجی را با همان «نام فایل» ذخیره کنید."),
 ("کار من", "تمام متن‌ها، ترجمه‌ها، جمله‌های مثال، تمرین‌ها و پرامپت‌ها را نوشته‌ام و ادامه می‌دهم."),
 ("", ""),
 ("مهم — سبک یکدست", "پرامپت‌ها همگی با یک پیشوند سبک ثابت شروع می‌شوند. آن را تغییر ندهید تا تمام تصاویر اپ یکدست به‌نظر برسند."),
 ("مهم — بدون متن", "در پرامپت‌ها «no text» آمده. اگر مدل باز هم حرف انگلیسی داخل تصویر گذاشت، دوباره بسازید."),
 ("مهم — نام فایل", "نام فایل را دقیقاً همان‌طور که در ستون آمده بگذارید. کد اپ بر اساس همین نام تصویر را پیدا می‌کند."),
 ("", ""),
 ("صوت", "نیازی به ضبط صدا نیست. ستون «متن صوتی» ورودی موتور TTS است؛ فایل‌های صوتی را یک‌بار روی سرور می‌سازیم."),
 ("", ""),
 ("وضعیت پیشرفت", "ستون «وضعیت» را بعد از ساخت هر تصویر روی «آماده» بگذارید تا بدانیم کجا هستیم."),
 ("", ""),
 ("قاعده ۶۰٪ تولیدی", "طبق سند ۰۷، حداقل ۶۰٪ تمرین‌ها باید «تولیدی» باشند (کاربر خودش می‌سازد) نه «تشخیصی» (انتخاب از گزینه). ستون «دسته» در برگه تمرین‌ها همین را نشان می‌دهد."),
 ("ستون واژه هدف", "برای لایتنر خودکار است: اگر کاربر تمرین را غلط جواب دهد، این واژه خودکار وارد جعبه لایتنر او می‌شود."),
 ("پاسخ‌های جایگزین", "شکل‌های درست دیگر همان پاسخ. کاربر نباید بابت تفاوت جزئی (مثلاً I'm به‌جای I am) جریمه شود."),
 ("", ""),
 ("برگه‌ها", "نقشه دروس · واژگان · مکالمه · گرامر · تمرین‌ها"),
]
for a, b in rows: g.append([a, b])
g["A1"].font = Font(bold=True, size=16, color="2F3E7E")
for r in range(3, g.max_row + 1):
    g.cell(r, 1).font = Font(bold=True)
    g.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="center")
    g.row_dimensions[r].height = 30

# --- نقشه دروس
ws = sheet(wb, "نقشه دروس", ["درس","عنوان انگلیسی","موضوع گرامر","ماموریت درس","تعداد واژه","رایگان؟","وضعیت"],
           [8, 26, 34, 28, 12, 10, 12])
for n, t, gr, th in SYLLABUS:
    ws.append([n, t, gr, th, 20, "بله" if n <= 5 else "خیر", "آماده" if n in V else "در انتظار"])
finish(ws, 7, wrap_cols=(3, 4))

# --- واژگان
ws = sheet(wb, "واژگان",
    ["شناسه","درس","واژه","بخش کلام","تلفظ IPA","معنی فارسی","جمله مثال","ترجمه جمله",
     "پرامپت تصویر (کپی کنید)","نام فایل تصویر","متن صوتی","وضعیت"],
    [16, 7, 18, 13, 17, 20, 34, 30, 78, 30, 34, 11])
# دروس ۱۱ به بعد در فایل جدا هستند تا این فایل قابل خواندن بماند
from data_a1_11_20 import V as V_11_20
from data_a1_21_30 import V as V_21_30
V.update(V_11_20); V.update(V_21_30)

for ln in sorted(V):
    for i, (w, pos, ipa, fa, ex, exfa, scene) in enumerate(V[ln], 1):
        slug = w.lower().replace(" ", "_").replace("'", "")
        ws.append([f"A1-L{ln:02d}-W{i:02d}", ln, w, pos, ipa, fa, ex, exfa,
                   STYLE + scene, f"a1_l{ln:02d}_w{i:02d}_{slug}.png", f"{w}. {ex}", "در انتظار تصویر"])
finish(ws, 12, wrap_cols=(7, 8, 9, 11))
for r in range(2, ws.max_row + 1):
    ws.cell(r, 9).fill = TODO
    ws.row_dimensions[r].height = 46

# --- مکالمه
ws = sheet(wb, "مکالمه", ["درس","نوبت","گوینده","جمله انگلیسی","ترجمه فارسی","صدای TTS","وضعیت"],
           [8, 8, 12, 42, 38, 14, 12])
DLG = {
 1: [("A","Hello! My name is Sara. What's your name?","سلام! اسم من ساراست. اسم شما چیست؟"),
     ("B","Hi Sara. I'm Ali. Nice to meet you.","سلام سارا. من علی هستم. از ملاقاتت خوشبختم."),
     ("A","Nice to meet you too. Are you a student?","من هم همینطور. تو دانش‌آموزی؟"),
     ("B","No, I'm a teacher. And you?","نه، من معلم هستم. تو چطور؟"),
     ("A","I'm a student. Goodbye, Ali!","من دانش‌آموزم. خداحافظ علی!"),
     ("B","Goodbye, Sara. See you tomorrow.","خداحافظ سارا. فردا می‌بینمت.")],
 2: [("A","What is this?","این چیست؟"),
     ("B","It's an umbrella. It's new.","این یک چتر است. نو است."),
     ("A","And what is that on the table?","و آن روی میز چیست؟"),
     ("B","That's my book and a blue pen.","آن کتاب من و یک خودکار آبی است."),
     ("A","Is this your bag?","این کیف توست؟"),
     ("B","Yes, it is. Thanks!","بله، هست. ممنون!")],
 4: [("A","Who is this in the photo?","این در عکس کیست؟"),
     ("B","This is my sister. Her name is Maryam.","این خواهر من است. اسمش مریم است."),
     ("A","Is she a student?","او دانش‌آموز است؟"),
     ("B","No, she's a nurse. And that's my father.","نه، او پرستار است. و آن پدر من است."),
     ("A","Your family is very nice.","خانواده‌ات خیلی خوب است."),
     ("B","Thank you!","ممنون!")],
 8: [("A","What do you do every day?","هر روز چه کار می‌کنی؟"),
     ("B","I wake up at six and I study English.","ساعت شش بیدار می‌شوم و انگلیسی می‌خوانم."),
     ("A","Do you go to school?","به مدرسه می‌روی؟"),
     ("B","Yes, I go to school at eight.","بله، ساعت هشت به مدرسه می‌روم."),
     ("A","And at night?","و شب‌ها؟"),
     ("B","I read a book and sleep at ten.","کتاب می‌خوانم و ساعت ده می‌خوابم.")],
 10:[("A","Excuse me, what time is it?","ببخشید، ساعت چند است؟"),
     ("B","It's half past two.","دو و نیم است."),
     ("A","Thanks. Is the class on Monday?","ممنون. کلاس دوشنبه است؟"),
     ("B","No, it's on Wednesday at four o'clock.","نه، چهارشنبه ساعت چهار است."),
     ("A","Okay. See you tomorrow!","باشه. فردا می‌بینمت!"),
     ("B","Goodbye!","خداحافظ!")],
 3: [("A","Who is that man?","آن مرد کیست؟"),
     ("B","He is my uncle. He is a doctor.","او عموی من است. پزشک است."),
     ("A","Is he tall?","او قدبلند است؟"),
     ("B","Yes, he is very tall.","بله، خیلی قدبلند است."),
     ("A","And who is that girl?","و آن دختر کیست؟"),
     ("B","She is his daughter. She is happy today.","او دخترش است. امروز خوشحال است.")],
 6: [("A","Where are you from?","اهل کجایی؟"),
     ("B","I am from Iran. And you?","من اهل ایرانم. تو چطور؟"),
     ("A","I am from a small city near Shiraz.","من اهل شهر کوچکی نزدیک شیرازم."),
     ("B","Is it beautiful?","زیباست؟"),
     ("A","Yes, very beautiful. Where do you live now?","بله، خیلی زیبا. الان کجا زندگی می‌کنی؟"),
     ("B","I live in Tehran. It is a big city.","در تهران زندگی می‌کنم. شهر بزرگی است.")],
 9: [("A","What does your brother do?","برادرت چه کار می‌کند؟"),
     ("B","He works in an office. He always starts at eight.","در یک دفتر کار می‌کند. همیشه ساعت هشت شروع می‌کند."),
     ("A","Does he study too?","درس هم می‌خواند؟"),
     ("B","Yes, he studies English every evening.","بله، هر عصر انگلیسی می‌خواند."),
     ("A","He is very busy!","خیلی مشغول است!"),
     ("B","Yes, but he never forgets his family.","بله، ولی هرگز خانواده‌اش را فراموش نمی‌کند.")],
}
from data_a1_11_20_b import DLG as DLG_11_20
from data_a1_21_30_b import DLG as DLG_21_30
DLG.update(DLG_11_20); DLG.update(DLG_21_30)

for ln in sorted(DLG):
    for i, (spk, en, fa) in enumerate(DLG[ln], 1):
        ws.append([ln, i, "نفر اول" if spk == "A" else "نفر دوم", en, fa,
                   "زن" if spk == "A" else "مرد", "آماده"])
finish(ws, 7, wrap_cols=(4, 5))

# --- گرامر
ws = sheet(wb, "گرامر",
           ["درس","کارت","عنوان","توضیح فارسی","مثال‌ها",
            "پرسش بررسی","پاسخ درست","گزینه‌ها","نکته هر گزینه غلط","وضعیت"],
           [8, 8, 30, 62, 46, 34, 20, 30, 60, 12])
GR = {
 1: [("افعال to be",
      "فعل «بودن» در فارسی «است / هستم / هستند» است؛ در انگلیسی سه شکل دارد و شکلش به فاعل جمله بستگی دارد.\n"
      "با I همیشه am می‌آید.\n"
      "با he، she و it همیشه is می‌آید.\n"
      "با you، we و they همیشه are می‌آید.\n"
      "برخلاف فارسی، این فعل هیچ‌وقت از جمله حذف نمی‌شود.",
      "I am a student.\nHe is a manager.\nShe is a teacher.\nIt is a book.\nYou are my friend.\nWe are here.\nThey are teachers."),
     ("ضمائر فاعلی",
      "ضمیر فاعلی جای اسم می‌نشیند تا مجبور نباشیم اسم را تکرار کنیم.\n"
      "I (من) · you (تو یا شما) · he (او، برای آقایان) · she (او، برای خانم‌ها)\n"
      "it (آن، برای اشیا و حیوانات) · we (ما) · they (آنها، هم افراد هم اشیا)\n"
      "ضمیر فاعلی همیشه اول جمله می‌آید.\n"
      "حرف I در هر جای جمله با حرف بزرگ نوشته می‌شود.",
      "Ali is a doctor. → He is a doctor.\nSara is a teacher. → She is a teacher.\nThe book is new. → It is new.\nAli and I are friends. → We are friends.")],
 2: [("a و an",
      "پیش از اسم مفرد قابل شمارش می‌آید و معنی «یک» می‌دهد.\n"
      "a پیش از واژه‌هایی که با صدای بی‌صدا شروع می‌شوند.\n"
      "an پیش از واژه‌هایی که با صدای صدادار شروع می‌شوند: a, e, i, o, u\n"
      "ملاک «صدا»ست نه حرف: an hour چون h خوانده نمی‌شود.\n"
      "پیش از اسم جمع هرگز a یا an نمی‌آید.",
      "a book · a pen · a car\nan apple · an umbrella · an office\nan hour (نه a hour)\nbooks (نه a books)"),
     ("اسم‌های جمع",
      "برای جمع بستن معمولاً s به آخر اسم اضافه می‌شود.\n"
      "اگر اسم به s, x, ch, sh یا o ختم شود، es می‌گیرد.\n"
      "اگر به حرف بی‌صدا + y ختم شود، y به ies تبدیل می‌شود.\n"
      "بعضی اسم‌ها جمع بی‌قاعده دارند و باید حفظ شوند.",
      "book → books · pen → pens\nbox → boxes · watch → watches\ncity → cities · baby → babies\nman → men · woman → women · child → children")],
 3: [("سوالی کردن با to be",
      "جای فعل و فاعل عوض می‌شود و فعل به ابتدای جمله می‌رود.\n"
      "برای پاسخ کوتاه، فعل تکرار می‌شود نه کل جمله.\n"
      "پاسخ مثبت: Yes, I am. — پاسخ منفی: No, he isn't.",
      "You are a teacher. → Are you a teacher?\nHe is happy. → Is he happy?\nThey are students. → Are they students?\nIs she a doctor? → Yes, she is.\nAre you tired? → No, I'm not."),
     ("who و what و where",
      "who برای پرسش درباره افراد به‌کار می‌رود.\n"
      "what برای اشیا و مفاهیم.\n"
      "where برای مکان.\n"
      "کلمه پرسشی همیشه اول جمله می‌آید، بعد فعل، بعد فاعل.",
      "Who is he? → He is my brother.\nWhat is this? → It's a pen.\nWhere is my bag? → It's on the table.\nWho are they? → They are my friends.")],
 4: [("صفات ملکی",
      "پیش از اسم می‌آید و نشان می‌دهد چیزی مال کیست.\n"
      "my (مالِ من) · your (مالِ تو یا شما) · his (مالِ او، آقا) · her (مالِ او، خانم)\n"
      "its (مالِ آن) · our (مالِ ما) · their (مالِ آنها)\n"
      "برخلاف فارسی، پیش از اسم می‌آید نه بعد از آن.",
      "This is my book.\nHer name is Sara.\nOur school is big.\nTheir car is new.\nHis father is a doctor."),
     ("'s مالکیت",
      "برای نشان دادن مالکیت افراد، به آخر اسم 's اضافه می‌شود.\n"
      "اگر اسم جمع باشد و به s ختم شود، فقط ' اضافه می‌شود.\n"
      "ترتیب برعکس فارسی است: اول مالک، بعد شیء.",
      "Ali's car — ماشینِ علی\nmy sister's bag — کیفِ خواهرم\nthe students' books — کتاب‌های دانش‌آموزان\nSara's phone is new.")],
 5: [("پرسش سن با how old",
      "ساختار: how old + فعل to be + فاعل\n"
      "در پاسخ می‌توان years old گفت یا فقط عدد.\n"
      "برخلاف فارسی که «چند سال داری» می‌گوییم، انگلیسی «چقدر پیر هستی» می‌گوید.",
      "How old are you? → I am twenty years old.\nHow old is he? → He is ten.\nHow old is your sister? → She's twenty-five."),
     ("اعداد ترتیبی",
      "اعداد اصلی برای شمارش به‌کار می‌روند و ترتیبی برای رتبه.\n"
      "بیشتر اعداد ترتیبی پسوند th می‌گیرند.\n"
      "سه استثنا وجود دارد که باید حفظ شوند: first, second, third",
      "one → first · two → second · three → third\nfour → fourth · five → fifth · six → sixth\nThis is my first day.\nHe lives on the third floor.")],
 6: [("پرسش با where و from",
      "برای پرسیدن اهلیت یا محل زندگی به‌کار می‌رود.\n"
      "پاسخ با from به‌همراه نام کشور یا شهر می‌آید.\n"
      "کلمه پرسشی اول جمله، بعد فعل to be، بعد فاعل.",
      "Where are you from? → I am from Iran.\nWhere is she from? → She is from Tehran.\nWhere are they from? → They are from Shiraz."),
     ("فعل live و حروف اضافه مکان",
      "با فعل live از حرف اضافه in برای شهر و کشور استفاده می‌شود.\n"
      "near یعنی نزدیک و far یعنی دور.\n"
      "برای ساختمان و نشانی دقیق از at استفاده می‌شود.",
      "I live in Tehran.\nMy school is near my house.\nThe airport is far from here.\nShe lives in a small village.")],
 7: [("مشاغل با a و an",
      "پیش از نام شغل حتماً a یا an می‌آید — این نکته برخلاف فارسی است.\n"
      "در فارسی می‌گوییم «او پرستار است»، اما انگلیسی «او یک پرستار است».\n"
      "اگر شغل با صدای صدادار شروع شود، an می‌گیرد.",
      "She is a nurse.\nHe is an engineer.\nI am a student.\nMy father is a driver.\nThey are teachers. (جمع، بدون a)"),
     ("پرسش شغل",
      "دو راه رایج وجود دارد: What is your job? یا What do you do?\n"
      "شکل دوم طبیعی‌تر و رایج‌تر است.\n"
      "برای سوم‌شخص از does استفاده می‌شود.",
      "What is your job? → I am a teacher.\nWhat do you do? → I work in an office.\nWhat does he do? → He is a driver.")],
 8: [("حال ساده — I / you / we / they",
      "برای عادت‌ها، کارهای همیشگی و حقایق کلی به‌کار می‌رود.\n"
      "با این فاعل‌ها شکل فعل هیچ تغییری نمی‌کند.\n"
      "قیدهای زمان مثل every day معمولاً آخر جمله می‌آیند.",
      "I go to school.\nThey play football.\nWe study English every day.\nYou work in a shop."),
     ("منفی و سوالی در حال ساده",
      "برای منفی کردن از don't پیش از فعل استفاده می‌شود.\n"
      "برای سوالی کردن Do به ابتدای جمله می‌آید.\n"
      "در جمله منفی و سوالی، خود فعل بدون تغییر می‌ماند.\n"
      "پاسخ کوتاه با do یا don't داده می‌شود.",
      "I don't like tea.\nThey don't work here.\nDo you play football? → Yes, I do.\nDo they live in Tehran? → No, they don't.")],
 9: [("سوم‌شخص مفرد",
      "با he، she و it فعل s می‌گیرد — این مهم‌ترین نکته حال ساده است.\n"
      "اگر فعل به حرف بی‌صدا + y ختم شود، y به ies تبدیل می‌شود.\n"
      "چند فعل بی‌قاعده‌اند: go → goes، do → does، have → has\n"
      "در جمله منفی و سوالی، doesn't و Does می‌آید و فعل s نمی‌گیرد.",
      "He works every day.\nShe studies English.\nIt goes fast.\nHe has a car.\nShe doesn't work here. (نه doesn't works)"),
     ("قیدهای تکرار",
      "نشان می‌دهند کاری هر چند وقت یک‌بار انجام می‌شود.\n"
      "always (همیشه) · usually (معمولاً) · often (اغلب) · sometimes (گاهی) · never (هرگز)\n"
      "پیش از فعل اصلی می‌آیند، اما بعد از فعل to be.\n"
      "با never فعل مثبت می‌ماند و دوباره منفی نمی‌شود.",
      "He always drinks tea.\nI usually walk to work.\nShe is always happy. (بعد از is)\nI never eat meat. (نه never don't eat)")],
 10:[("حروف اضافه زمان",
      "at برای ساعت دقیق به‌کار می‌رود.\n"
      "on برای روز هفته و تاریخ.\n"
      "in برای ماه، سال، فصل و بخش‌های روز.\n"
      "سه استثنا: at night و at noon و at midnight",
      "at three o'clock · at night\non Monday · on Friday\nin the morning · in July · in winter\nThe class is on Wednesday at four."),
     ("پرسش و گفتن ساعت",
      "دو شکل رایج: What time is it? یا What's the time?\n"
      "ساعت کامل با o'clock بیان می‌شود.\n"
      "نیم ساعت با half past و ربع با quarter past یا quarter to.\n"
      "در انگلیسی محاوره‌ای، دقیقه اول گفته می‌شود و بعد ساعت.",
      "It's three o'clock.\nIt's half past two. — دو و نیم\nIt's quarter past five. — پنج و ربع\nIt's quarter to nine. — یک ربع به نه")],
}
from data_a1_11_20_b import GR as GR_11_20
from data_a1_21_30_b import GR as GR_21_30
GR.update(GR_11_20); GR.update(GR_21_30)
from data_a1_checks import CHECK

for ln in sorted(GR):
    for i, (t, d, ex) in enumerate(GR[ln], 1):
        chk = CHECK.get((ln, i))
        if chk is None:
            raise SystemExit(
                f"درس {ln} کارت گرامر {i} پرسش بررسی ندارد. "
                f"هر کارت گرامر باید یک پرسش داشته باشد، وگرنه بخش آموزش "
                f"دوباره فقط خواندن می‌شود."
            )
        q, ans, opts, tips = chk
        if ans not in opts:
            raise SystemExit(f"درس {ln} کارت {i}: پاسخ درست بین گزینه‌ها نیست.")
        wrong = [o for o in opts if o != ans]
        missing = [o for o in wrong if not tips.get(o)]
        if missing:
            raise SystemExit(
                f"درس {ln} کارت {i}: گزینه‌های {missing} نکته ندارند. "
                f"نکته هدفمند روی خطا اصلِ این بخش است — «اشتباه است» کافی نیست."
            )
        ws.append([
            ln, f"Grammar {i}", t, d, ex,
            q, ans, " | ".join(opts),
            "\n".join(f"{o} ← {tips[o]}" for o in wrong),
            "آماده",
        ])
finish(ws, 10, wrap_cols=(4, 5, 6, 9))
for r in range(2, ws.max_row + 1): ws.row_dimensions[r].height = 62

# --- تمرین‌ها  (قاعده: حداقل ۶۰٪ تولیدی — سند ۰۷ تمایز ۱)
ws = sheet(wb, "تمرین‌ها",
    ["درس","تمرین","دسته","نوع","صورت سوال","پاسخ درست","پاسخ‌های جایگزین",
     "گزینه‌ها","واژه هدف (لایتنر)","راهنما","وضعیت"],
    [7, 11, 11, 17, 40, 26, 26, 32, 18, 32, 10])
P, R = "تولیدی", "تشخیصی"
# (دسته, نوع, صورت سوال, پاسخ, جایگزین‌ها, گزینه‌ها, واژه هدف, راهنما)
EX = {
 1: [(R,"چندگزینه‌ای","I ___ a student.","am","","am | is | are","","با I همیشه am می‌آید."),
     (R,"تطبیق","teacher","معلم","","teacher | student | doctor","teacher","واژه را به معنی درست وصل کن."),
     (P,"جای خالی","He ___ a manager.","is","","","manager","با he/she/it از is استفاده کن."),
     (P,"متن آزاد","با شکل مخفف بنویسید: I am Ali.","I'm Ali.","I'm Ali","","","I am → I'm"),
     (P,"متن آزاد","با شکل مخفف بنویسید: She is a teacher.","She's a teacher.","She's a teacher","","teacher","She is → She's"),
     (P,"مرتب‌سازی","is / my / this / friend","This is my friend.","","","friend","جمله با فاعل شروع می‌شود."),
     (P,"ترجمه به انگلیسی","او یک پزشک است.","He is a doctor.","He's a doctor.","","doctor","با he فعل is می‌آید."),
     (P,"گفتار","Nice to meet you.","Nice to meet you.","","","meet","عبارت را واضح تلفظ کنید.")],
 2: [(R,"چندگزینه‌ای","It's ___ umbrella.","an","","a | an | the","umbrella","umbrella با صدای صدادار شروع می‌شود."),
     (R,"تطبیق","key","کلید","","key | door | window","key","واژه را به معنی درست وصل کن."),
     (P,"جای خالی","This is ___ book.","a","","","book","book با صدای بی‌صدا شروع می‌شود."),
     (P,"متن آزاد","جمع ببندید: box","boxes","","","box","اسم‌های مختوم به x در جمع es می‌گیرند."),
     (P,"متن آزاد","جمع ببندید: watch","watches","","","watch","اسم‌های مختوم به ch در جمع es می‌گیرند."),
     (P,"مرتب‌سازی","an / is / it / umbrella","It is an umbrella.","It's an umbrella.","","umbrella","با ضمیر شروع کن."),
     (P,"ترجمه به انگلیسی","این یک کتاب است.","This is a book.","","","book","برای اشاره نزدیک this می‌آید."),
     (P,"گفتار","It's an umbrella.","It's an umbrella.","","","umbrella","عبارت را واضح تلفظ کنید.")],
 3: [(R,"چندگزینه‌ای","___ is he? — He is my brother.","Who","","Who | What | Where","who","برای پرسش درباره افراد who می‌آید."),
     (R,"تطبیق","happy","خوشحال","","happy | sad | old","happy","واژه را به معنی درست وصل کن."),
     (P,"جای خالی","___ is this? — It's a pen.","What","","","what","برای اشیا what می‌آید."),
     (P,"متن آزاد","سوالی کنید: You are a teacher.","Are you a teacher?","","","","جای فعل و فاعل را عوض کن."),
     (P,"متن آزاد","سوالی کنید: She is happy.","Is she happy?","","","happy","فعل به ابتدای جمله می‌رود."),
     (P,"مرتب‌سازی","he / is / very / tall","He is very tall.","He's very tall.","","tall","صفت بعد از فعل to be می‌آید."),
     (P,"ترجمه به انگلیسی","او کیست؟","Who is he?","Who's he?","","who","برای افراد who به‌کار می‌رود."),
     (P,"گفتار","Who is that woman?","Who is that woman?","","","woman","عبارت را واضح تلفظ کنید.")],
 4: [(R,"چندگزینه‌ای","This is ___ book. (مالِ من)","my","","my | your | his","my","برای «مالِ من» از my استفاده کن."),
     (R,"تطبیق","sister","خواهر","","sister | brother | mother","sister","واژه را به معنی درست وصل کن."),
     (P,"جای خالی","___ name is Sara. (او، زن)","Her","","","","برای مؤنث her می‌آید."),
     (P,"متن آزاد","با 's بنویسید: the car of Ali","Ali's car","","","car","به آخر اسم شخص 's اضافه کن."),
     (P,"متن آزاد","با 's بنویسید: the bag of my sister","my sister's bag","","","sister","به آخر اسم 's اضافه کن."),
     (P,"مرتب‌سازی","my / this / sister / is","This is my sister.","","","sister","با اشاره‌گر شروع کن."),
     (P,"ترجمه به انگلیسی","مدرسه ما بزرگ است.","Our school is big.","","","our","برای «مالِ ما» our می‌آید."),
     (P,"گفتار","This is my family.","This is my family.","","","family","عبارت را واضح تلفظ کنید.")],
 5: [(R,"چندگزینه‌ای","How ___ are you?","old","","old | many | much","old","برای پرسش سن how old می‌آید."),
     (R,"تطبیق","birthday","تولد","","birthday | party | gift","birthday","واژه را به معنی درست وصل کن."),
     (P,"جای خالی","I am twenty years ___.","old","","","age","بعد از سن واژه old می‌آید."),
     (P,"متن آزاد","با حروف بنویسید: 15","fifteen","","","number","اعداد ۱۳ تا ۱۹ پسوند teen می‌گیرند."),
     (P,"متن آزاد","با حروف بنویسید: 30","thirty","","","number","اعداد دهگان پسوند ty می‌گیرند."),
     (P,"مرتب‌سازی","old / how / you / are","How old are you?","","","age","جمله پرسشی با how شروع می‌شود."),
     (P,"ترجمه به انگلیسی","من بیست ساله هستم.","I am twenty years old.","I'm twenty years old.","","year","ساختار: فاعل + be + عدد + years old"),
     (P,"گفتار","How old are you?","How old are you?","","","age","عبارت را واضح تلفظ کنید.")],
 6: [(R,"چندگزینه‌ای","I am ___ Iran.","from","","from | in | at","from","برای اهلیت from می‌آید."),
     (R,"تطبیق","city","شهر","","city | country | village","city","واژه را به معنی درست وصل کن."),
     (P,"جای خالی","___ are you from?","Where","","","where","برای پرسش مکان where می‌آید."),
     (P,"متن آزاد","سوالی کنید: She is from Tehran.","Is she from Tehran?","","","from","فعل به ابتدای جمله می‌رود."),
     (P,"متن آزاد","کامل کنید با live: We ___ in Shiraz.","live","","","live","با we فعل s نمی‌گیرد."),
     (P,"مرتب‌سازی","from / where / you / are","Where are you from?","","","where","با کلمه پرسشی شروع کن."),
     (P,"ترجمه به انگلیسی","من اهل ایران هستم.","I am from Iran.","I'm from Iran.","","country","ساختار: فاعل + be + from + کشور"),
     (P,"گفتار","Where are you from?","Where are you from?","","","where","عبارت را واضح تلفظ کنید.")],
 7: [(R,"چندگزینه‌ای","She is ___ engineer.","an","","a | an | the","engineer","engineer با صدای صدادار شروع می‌شود."),
     (R,"تطبیق","nurse","پرستار","","nurse | doctor | cook","nurse","واژه را به معنی درست وصل کن."),
     (P,"جای خالی","My father is ___ driver.","a","","","driver","پیش از نام شغل a یا an می‌آید."),
     (P,"متن آزاد","سوالی کنید: He is a cook.","Is he a cook?","","","cook","فعل به ابتدای جمله می‌رود."),
     (P,"متن آزاد","منفی کنید: She is a nurse.","She is not a nurse.","She isn't a nurse.","","nurse","not بعد از فعل to be می‌آید."),
     (P,"مرتب‌سازی","job / what / your / is","What is your job?","What's your job?","","job","با کلمه پرسشی شروع کن."),
     (P,"ترجمه به انگلیسی","او مهندس است.","He is an engineer.","He's an engineer.","","engineer","پیش از شغل حتماً a/an بگذار."),
     (P,"گفتار","What is your job?","What is your job?","","","job","عبارت را واضح تلفظ کنید.")],
 8: [(R,"چندگزینه‌ای","I ___ to school every day.","go","","go | goes | going","go","با I فعل s نمی‌گیرد."),
     (R,"تطبیق","study","درس خواندن","","study | play | read","study","واژه را به معنی درست وصل کن."),
     (P,"جای خالی","They ___ football. (play)","play","","","play","با they فعل s نمی‌گیرد."),
     (P,"متن آزاد","منفی کنید: I like tea.","I don't like tea.","I do not like tea.","","drink","منفی حال ساده با don't ساخته می‌شود."),
     (P,"متن آزاد","سوالی کنید: You play football.","Do you play football?","","","play","سوال حال ساده با Do شروع می‌شود."),
     (P,"مرتب‌سازی","school / to / go / I","I go to school.","","","go","با فاعل شروع کن."),
     (P,"ترجمه به انگلیسی","ما انگلیسی می‌خوانیم.","We study English.","","","study","با we فعل بدون s می‌آید."),
     (P,"گفتار","I wake up at six.","I wake up at six.","","","wake","عبارت را واضح تلفظ کنید.")],
 9: [(R,"چندگزینه‌ای","He ___ every day.","works","","work | works | working","work","با he فعل s می‌گیرد."),
     (R,"تطبیق","always","همیشه","","always | never | often","always","واژه را به معنی درست وصل کن."),
     (P,"جای خالی","She ___ English. (study)","studies","","","study","فعل مختوم به y با he/she به ies تبدیل می‌شود."),
     (P,"متن آزاد","با he بنویسید: I go to work.","He goes to work.","","","work","go با he به goes تبدیل می‌شود."),
     (P,"متن آزاد","قید را در جای درست بگذارید: He drinks tea. (always)","He always drinks tea.","","","always","قید تکرار پیش از فعل اصلی می‌آید."),
     (P,"مرتب‌سازی","usually / she / walks / to work","She usually walks to work.","","","usually","قید تکرار بعد از فاعل می‌آید."),
     (P,"ترجمه به انگلیسی","او هرگز گوشت نمی‌خورد.","He never eats meat.","","","never","با never فعل مثبت می‌ماند."),
     (P,"گفتار","He works every day.","He works every day.","","","work","عبارت را واضح تلفظ کنید.")],
 10:[(R,"چندگزینه‌ای","The class is ___ Monday.","on","","at | on | in","Monday","برای روزهای هفته on می‌آید."),
     (R,"تطبیق","clock","ساعت دیواری","","clock | watch | time","clock","واژه را به معنی درست وصل کن."),
     (P,"جای خالی","I study ___ the morning.","in","","","morning","برای بخش‌های روز in می‌آید."),
     (P,"متن آزاد","کامل کنید: It's three ___.","o'clock","","","o'clock","برای ساعت کامل o'clock می‌آید."),
     (P,"متن آزاد","به انگلیسی بنویسید: ۲:۳۰","It's half past two.","half past two","","half","نیم ساعت با half past بیان می‌شود."),
     (P,"مرتب‌سازی","is / what / it / time","What time is it?","","","time","با کلمه پرسشی شروع کن."),
     (P,"ترجمه به انگلیسی","ساعت چند است؟","What time is it?","What's the time?","","time","دو شکل درست دارد."),
     (P,"گفتار","It's half past two.","It's half past two.","","","half","عبارت را واضح تلفظ کنید.")],
}
from data_a1_11_20_b import EX as EX_11_20
from data_a1_21_30_b import EX as EX_21_30
EX.update(EX_11_20); EX.update(EX_21_30)


# تمرین «تطبیق» صورت سوالش واژه انگلیسی و پاسخش معنی فارسی است، پس
# گزینه‌ها هم باید فارسی باشند. گزینه‌ها را انگلیسی نوشته بودیم و نتیجه
# تمرینی بود که هیچ‌وقت جواب درست نداشت — کاربر هرچه می‌زد غلط می‌شد.
# ترجمه از خودِ جدول واژگان می‌آید تا دو جا از هم جدا نیفتند.
MEANING = {w.lower(): fa for rows in V.values() for (w, _p, _i, fa, *_r) in rows}

def translate_options(opt, answer, lesson, index):
    words = [x.strip() for x in opt.split("|") if x.strip()]
    out = []
    for w in words:
        fa = MEANING.get(w.lower())
        if fa is None:
            raise SystemExit(
                f"درس {lesson} تمرین {index}: معنی «{w}» در جدول واژگان نیست. "
                f"یا واژه را اضافه کنید یا گزینه را عوض کنید."
            )
        out.append(fa)
    if answer not in out:
        raise SystemExit(
            f"درس {lesson} تمرین {index}: پاسخ درست «{answer}» بین گزینه‌ها نیست "
            f"({' | '.join(out)}) — تمرین جواب‌ناپذیر می‌شود."
        )
    return " | ".join(out)

for ln in sorted(EX):
    for i, (cat, ty, q, a, alt, opt, tw, h) in enumerate(EX[ln], 1):
        if ty == "تطبیق" and opt:
            opt = translate_options(opt, a, ln, i)
        ws.append([ln, f"Exercise {i}", cat, ty, q, a, alt, opt, tw, h, "آماده"])
finish(ws, 11, wrap_cols=(5, 7, 8, 10))


# --- داستانک (اختیاری — فقط اگر stories_<level>.py وجود داشته باشد)
try:
    STORIES = __import__("stories_a1").STORIES
except ImportError:
    STORIES = {}

if STORIES:
    ws = sheet(wb, "داستانک",
               ["درس", "عنوان", "نوع", "متن انگلیسی", "ترجمه فارسی",
                "پاسخ درست", "گزینه‌ها", "راهنما", "وضعیت"],
               [7, 26, 9, 60, 46, 22, 34, 30, 10])
    for _ln in sorted(STORIES):
        _title, _paras, _qs = STORIES[_ln]
        for _en, _fa in _paras:
            ws.append([_ln, _title, "بند", _en, _fa, "", "", "", "آماده"])
        for _q, _qfa, _ans, _opts, _hint in _qs:
            ws.append([_ln, _title, "پرسش", _q, _qfa, _ans, " | ".join(_opts), _hint, "آماده"])
    finish(ws, 9, wrap_cols=(4, 5, 7, 8))

wb.save(OUT)

nv = sum(len(x) for x in V.values())
nd = sum(len(x) for x in DLG.values())
ng = sum(len(x) for x in GR.values())
ne = sum(len(x) for x in EX.values())
np_ = sum(1 for x in EX.values() for e in x if e[0] == P)
print(f"ساخته شد: {OUT}")
print(f"  واژگان : {nv} ردیف  ({len(V)} درس)  ← {nv} تصویر لازم است")
print(f"  مکالمه : {nd} ردیف  ({len(DLG)} دیالوگ)")
print(f"  گرامر  : {ng} کارت  ({len(GR)} درس)")
print(f"  تمرین  : {ne} تمرین  ({len(EX)} درس)")
print(f"  سرفصل  : {len(SYLLABUS)} درس")
print()
print(f"  بررسی قاعده ۶۰٪ تولیدی: {np_}/{ne} = {np_/ne:.0%}", "✅" if np_/ne >= 0.60 else "❌")

# --- گزارش واژه‌های تکراری
# هر واژه که دو بار تدریس شود یک جای خالیِ هدررفته است: یک تصویر اضافه
# (با هزینه واقعی)، یک کارت لایتنر تکراری، و یک واژه تازه که جا نشد.
# A1 یک‌بار ۱۱۴ تکرار داشت (۴۸۶ واژه یکتا از ۶۰۰) و چون خطایی نمی‌داد
# کسی متوجهش نشد. پاک‌سازی شد و حالا مثل A2 ساخت را متوقف می‌کند.
_seen, _dups = {}, []
for _ln in sorted(V):
    for _row in V[_ln]:
        _k = _row[0].lower()
        if _k in _seen:
            _dups.append(f"{_row[0]} ({_seen[_k]}→{_ln})")
        else:
            _seen[_k] = _ln
if _dups:
    raise SystemExit(
        f"\n  ❌ واژه تکراری: {len(_dups)} از {nv} — فقط {len(_seen)} واژه یکتا\n"
        f"     {'، '.join(_dups)}"
    )
for ln in sorted(EX):
    p = sum(1 for e in EX[ln] if e[0] == P)
    flag = "✅" if p / len(EX[ln]) >= 0.60 else "❌"
    print(f"    درس {ln:2d}: {p}/{len(EX[ln])} تولیدی {flag}")
