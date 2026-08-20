#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
تولید کاربرگ محتوای سطح A2.

ساختار برگه‌ها دقیقاً همان A1 است تا `export_to_app.py` بدون هیچ تغییری
هر دو را بخواند. تنها تفاوت، منبع داده است:

    A1 → متن دروس ۱ تا ۱۰ داخل build_a1.py، بقیه در data_a1_*.py
    A2 → همه در data_a2_*.py و سرفصل در syllabus_a2.py

اجرا:
    python3 content/build_a2.py
    python3 content/export_to_app.py --level A2
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from syllabus_a2 import SYLLABUS

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "A2-محتوا.xlsx")

# --- محتوا: هر بلوک ده‌درسی یک فایل. دروسی که هنوز نوشته نشده‌اند
# فقط در سرفصل هستند و در خروجی نمی‌آیند — همان رفتار A1.
V, DLG, GR, EX = {}, {}, {}, {}
for mod in ("data_a2_01_10", "data_a2_04_10", "data_a2_11_20", "data_a2_21_30"):
    try:
        m = __import__(mod)
    except ImportError:
        continue
    V.update(getattr(m, "V", {}))
    DLG.update(getattr(m, "DLG", {}))
    GR.update(getattr(m, "GR", {}))
    EX.update(getattr(m, "EX", {}))
for mod in ("data_a2_01_10_b", "data_a2_04_10_b", "data_a2_11_20_b", "data_a2_21_30_b"):
    try:
        m = __import__(mod)
    except ImportError:
        continue
    DLG.update(getattr(m, "DLG", {}))
    GR.update(getattr(m, "GR", {}))
    EX.update(getattr(m, "EX", {}))

STYLE = ("Modern editorial illustration for a professional language-learning app. "
         "Scene: ")

HEAD_FILL = PatternFill("solid", fgColor="2F3E7E")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="D0D0D8")


def sheet(wb, title, headers, widths, rtl=True):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = rtl
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


def finish(ws, ncols, wrap_cols=()):
    for row in ws.iter_rows(min_row=2, max_col=ncols):
        for c in row:
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            c.alignment = Alignment(
                vertical="top",
                wrap_text=c.column in wrap_cols,
                horizontal="right",
            )


wb = Workbook()
wb.remove(wb.active)

# --- نقشه دروس
ws = sheet(wb, "نقشه دروس",
           ["درس", "عنوان انگلیسی", "موضوع گرامر", "ماموریت درس", "تعداد واژه", "رایگان؟", "وضعیت"],
           [7, 30, 30, 28, 11, 9, 10])
for n, title_en, grammar, theme, free in SYLLABUS:
    ws.append([n, title_en, grammar, theme, len(V.get(n, [])),
               "بله" if free else "خیر",
               "آماده" if n in V else "در انتظار محتوا"])
finish(ws, 7)

# --- محافظ تکرار واژه
# یک واژه که دو بار در دو درس تدریس شود یعنی یک جای خالیِ هدررفته: یک
# تصویر اضافه، یک کارت لایتنر تکراری و یک واژه تازه که جا نشد. اولین بار
# که این اتفاق افتاد ۳۸ واژه از ۴۰۰ هدر رفته بود و کسی متوجه نشد، چون
# خطایی رخ نمی‌داد. حالا ساخت متوقف می‌شود.
_seen = {}
_dups = []
for _ln in sorted(V):
    for _row in V[_ln]:
        _k = _row[0].lower()
        if _k in _seen:
            _dups.append(f"«{_row[0]}» در درس {_seen[_k]} و {_ln}")
        else:
            _seen[_k] = _ln
if _dups:
    raise SystemExit("واژه تکراری:\n  " + "\n  ".join(_dups))

# --- واژگان
ws = sheet(wb, "واژگان",
           ["شناسه", "درس", "واژه", "بخش کلام", "تلفظ IPA", "معنی فارسی", "جمله مثال", "ترجمه جمله",
            "پرامپت تصویر (کپی کنید)", "نام فایل تصویر", "متن صوتی", "وضعیت"],
           [16, 7, 18, 13, 17, 20, 34, 30, 78, 30, 34, 11])
for ln in sorted(V):
    for i, (w, pos, ipa, fa, ex, exfa, scene) in enumerate(V[ln], 1):
        slug = w.lower().replace(" ", "_").replace("'", "")
        ws.append([
            f"A2-L{ln:02d}-W{i:02d}", ln, w, pos, ipa, fa, ex, exfa,
            STYLE + scene,
            f"a2_l{ln:02d}_w{i:02d}_{slug}.png",
            f"{w}. {ex}",
            "آماده",
        ])
finish(ws, 12, wrap_cols=(7, 8, 9))

# --- مکالمه
ws = sheet(wb, "مکالمه",
           ["درس", "نوبت", "گوینده", "جمله انگلیسی", "ترجمه فارسی", "صدای TTS", "وضعیت"],
           [7, 8, 11, 42, 38, 11, 10])
for ln in sorted(DLG):
    for i, (spk, en, fa) in enumerate(DLG[ln], 1):
        ws.append([ln, i, "نفر اول" if spk == "A" else "نفر دوم", en, fa,
                   "زن" if spk == "A" else "مرد", "آماده"])
finish(ws, 7, wrap_cols=(4, 5))

# --- گرامر
ws = sheet(wb, "گرامر",
           ["درس", "کارت", "عنوان", "توضیح فارسی", "مثال‌ها",
            "پرسش بررسی", "پاسخ درست", "گزینه‌ها", "نکته هر گزینه غلط", "وضعیت"],
           [7, 9, 26, 60, 40, 34, 16, 26, 40, 10])
for ln in sorted(GR):
    for i, card in enumerate(GR[ln], 1):
        title, desc, examples = card[0], card[1], card[2]
        check_q = card[3] if len(card) > 3 else ""
        check_a = card[4] if len(card) > 4 else ""
        check_o = card[5] if len(card) > 5 else ""
        tips = card[6] if len(card) > 6 else ""
        ws.append([ln, f"Card {i}", title, desc, examples, check_q, check_a, check_o, tips, "آماده"])
finish(ws, 10, wrap_cols=(4, 5, 6, 8, 9))

# --- تمرین‌ها
ws = sheet(wb, "تمرین‌ها",
           ["درس", "تمرین", "دسته", "نوع", "صورت سوال", "پاسخ درست", "پاسخ‌های جایگزین",
            "گزینه‌ها", "واژه هدف (لایتنر)", "راهنما", "وضعیت"],
           [7, 11, 11, 17, 40, 26, 26, 32, 18, 32, 10])

P = "تولیدی"

# همان محافظ A1: گزینه‌های تمرین تطبیق باید فارسی باشند، وگرنه تمرین
# جواب درست ندارد. ترجمه از خود جدول واژگان می‌آید تا دو جا از هم جدا نیفتند.
MEANING = {w.lower(): fa for rows in V.values() for (w, _p, _i, fa, *_r) in rows}


def translate_options(opt, answer, lesson, index):
    words = [x.strip() for x in opt.split("|") if x.strip()]
    out = []
    for w in words:
        fa = MEANING.get(w.lower())
        if fa is None:
            raise SystemExit(
                f"درس {lesson} تمرین {index}: معنی «{w}» در جدول واژگان نیست."
            )
        out.append(fa)
    if answer not in out:
        raise SystemExit(
            f"درس {lesson} تمرین {index}: پاسخ درست «{answer}» بین گزینه‌ها نیست."
        )
    return " | ".join(out)


for ln in sorted(EX):
    for i, (cat, ty, q, a, alt, opt, tw, h) in enumerate(EX[ln], 1):
        if ty == "تطبیق" and opt:
            opt = translate_options(opt, a, ln, i)
        ws.append([ln, f"Exercise {i}", cat, ty, q, a, alt, opt, tw, h, "آماده"])
finish(ws, 11, wrap_cols=(5, 7, 8, 10))


# --- داستانک (اختیاری — فقط اگر stories_<level>.py وجود داشته باشد)
try:
    STORIES = __import__("stories_a2").STORIES
except ImportError:
    STORIES = {}

if STORIES:
    ws = sheet(wb, "داستانک",
               ["درس", "عنوان", "نوع", "متن انگلیسی", "ترجمه فارسی",
                "پاسخ درست", "گزینه‌ها", "راهنما", "وضعیت"],
               [7, 26, 9, 60, 46, 22, 34, 30, 10])
    for _ln in sorted(STORIES):
        _title, _paras, _qs = STORIES[_ln]
        for _en, _fa in _paras:
            ws.append([_ln, _title, "بند", _en, _fa, "", "", "", "آماده"])
        for _q, _qfa, _ans, _opts, _hint in _qs:
            ws.append([_ln, _title, "پرسش", _q, _qfa, _ans, " | ".join(_opts), _hint, "آماده"])
    finish(ws, 9, wrap_cols=(4, 5, 7, 8))

wb.save(OUT)

nv = sum(len(x) for x in V.values())
print(f"ساخته شد: {OUT}")
print(f"  سرفصل  : {len(SYLLABUS)} درس")
print(f"  واژگان : {nv} ردیف  ({len(V)} درس نوشته‌شده)  ← {nv} تصویر لازم است")
print(f"  مکالمه : {sum(len(x) for x in DLG.values())} ردیف")
print(f"  گرامر  : {sum(len(x) for x in GR.values())} کارت")
print(f"  تمرین  : {sum(len(x) for x in EX.values())} تمرین")

if EX:
    tp = sum(1 for rows in EX.values() for e in rows if e[0] == P)
    tt = sum(len(rows) for rows in EX.values())
    ok = "✅" if tp / tt >= 0.60 else "❌"
    print(f"\n  بررسی قاعده ۶۰٪ تولیدی: {tp}/{tt} = {tp*100//tt}% {ok}")

missing = [n for n, *_ in SYLLABUS if n not in V]
if missing:
    print(f"\n  ⏳ دروس بدون محتوا: {', '.join(map(str, missing))}")
