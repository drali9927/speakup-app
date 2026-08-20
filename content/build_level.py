#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
سازنده عمومی کاربرگ محتوا — برای هر سطحی که ساختار A2 را دنبال کند.

    python3 content/build_level.py --level B1
    python3 content/export_to_app.py --level B1

A1 و A2 سازنده خودشان را دارند (A1 چون ده درس اولش داخل خود فایل است).
از B1 به بعد همه از همین فایل استفاده می‌کنند تا محافظ‌ها یک‌جا بمانند:

  ۱. گزینه‌های تمرین تطبیق باید فارسی و شامل پاسخ درست باشند.
  ۲. هیچ واژه‌ای دو بار تدریس نشود.
  ۳. دو واژه‌ی یک درس معنی فارسیِ یکسان نگیرند.
  ۴. دست‌کم ۶۰٪ تمرین‌ها تولیدی باشند.
"""

import argparse
import importlib
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))

STYLE = ("Modern editorial illustration for a professional language-learning app. "
         "Scene: ")
HEAD_FILL = PatternFill("solid", fgColor="2F3E7E")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="D0D0D8")
P = "تولیدی"


def load(level):
    low = level.lower()
    syl = importlib.import_module(f"syllabus_{low}").SYLLABUS
    V, DLG, GR, EX, ST = {}, {}, {}, {}, {}

    # ماژول‌های داده **پیدا** می‌شوند و نه از فهرست ثابت خوانده.
    #
    # پیش‌تر فهرست ثابتِ ("01_10", "11_20", "21_30") بود. فایل
    # data_a2_04_10_b.py که بلوکش «04_10» است در آن فهرست نبود و
    # **هیچ‌وقت خوانده نمی‌شد** — نتیجه‌اش این بود که درس‌های ۴ تا ۱۰
    # سطح A2 بخش گرامرشان کاملاً خالی بود و هیچ خطایی هم نمی‌داد،
    # چون نبودِ کلید در دیکشنری خطا نیست.
    #
    # با پیدا کردن از روی نام فایل، افزودن هر بلوک تازه خودکار کار
    # می‌کند و این دسته اشکال دیگر ممکن نیست.
    names = sorted(
        f[:-3] for f in os.listdir(BASE)
        if f.startswith(f"data_{low}_") and f.endswith(".py")
    )
    for name in names:
        try:
            m = importlib.import_module(name)
        except ImportError:
            continue
        V.update(getattr(m, "V", {}))
        ST.update(getattr(m, "ST", {}))
        DLG.update(getattr(m, "DLG", {}))
        GR.update(getattr(m, "GR", {}))
        EX.update(getattr(m, "EX", {}))
    # داستانک‌ها در ماژول جدا: stories_<level>.py
    try:
        ST.update(importlib.import_module(f"stories_{low}").STORIES)
    except ImportError:
        pass
    return syl, V, DLG, GR, EX, ST


def visual_key(s):
    """شکل دیداریِ یک رشته فارسی — همتای normaliseForCompare در سمت اپ."""
    return re.sub(r"\s+", " ",
                  s.replace("‌", " ").replace("ي", "ی").replace("ك", "ک")).strip()


def gloss_key(s):
    """
    معنی پایه، بدون توضیح داخل پرانتز — همتای glossKey در سمت اپ.

    چرا لازم شد: «سلام» و «سلام (خودمانی)» دو رشته متفاوت‌اند و از
    visual_key رد می‌شدند، ولی برای زبان‌آموز یک چیزند. وقتی هر دو در
    گزینه‌های یک آزمون می‌آمدند، آزمون از سنجش به حدس‌زدن تبدیل می‌شد.
    این دقیقاً در درس اول A1 بود — نخستین آزمونی که هر کاربر می‌بیند.
    """
    return visual_key(re.sub(r"[(（][^)）]*[)）]", " ", s)).strip(" ،,-–")


def sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


def finish(ws, ncols, wrap_cols=()):
    for row in ws.iter_rows(min_row=2, max_col=ncols):
        for c in row:
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap_cols,
                                    horizontal="right")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--level", required=True)
    a = ap.parse_args()
    level = a.level.upper()
    low = level.lower()

    SYLLABUS, V, DLG, GR, EX, ST = load(level)
    out = os.path.join(BASE, f"{level}-محتوا.xlsx")

    # --- محافظ تکرار واژه
    seen, dups = {}, []
    for ln in sorted(V):
        for row in V[ln]:
            k = row[0].lower()
            if k in seen:
                dups.append(f"«{row[0]}» در درس {seen[k]} و {ln}")
            else:
                seen[k] = ln
    if dups:
        raise SystemExit("واژه تکراری:\n  " + "\n  ".join(dups))

    # --- محافظ معنیِ یکسان در یک درس
    #
    # دو سطح، چون دو چیز متفاوت‌اند:
    #
    # **خطا — معنیِ عیناً یکسان.** بازی جفت‌یابی برای هر دو کاشیِ فارسیِ
    # کاملاً یکسان می‌سازد و درست‌بودن را با pairKey می‌سنجد نه با متن؛
    # کاربر کاشی درست را می‌زند، غلط حساب می‌شود و واژه بی‌دلیل وارد
    # لایتنر می‌شود. این یک اشکال محتوایی است و باید ساخت را بخواباند.
    #
    # **هشدار — معنیِ پایه یکسان با پرانتز متفاوت.** مثل «که (برای
    # افراد)» و «که (برای اشیا)». این‌ها **درست‌اند**: who و which باید
    # در یک درس آموزش داده شوند و معنی فارسی‌شان هم واقعاً همین است.
    # مجبور کردن نویسنده به ساختن معنی مصنوعیِ متفاوت، محتوا را بدتر
    # می‌کند. محافظت واقعی سمت اپ است: `glossKey` نمی‌گذارد این دو با هم
    # در گزینه‌های یک آزمون بیایند. این‌جا فقط خبر می‌دهیم تا نویسنده
    # بداند.
    #
    # مقایسه روی شکل دیداری است، نه رشته خام: نیم‌فاصله و «ي/ك» عربی دو
    # رشته متفاوت می‌سازند که روی صفحه یکی دیده می‌شوند.
    same, near = [], []
    for ln in sorted(V):
        exact, base = {}, {}
        for row in V[ln]:
            ke = visual_key(row[3])
            if ke in exact:
                same.append(f"درس {ln}: «{exact[ke]}» و «{row[0]}» هر دو = «{row[3]}»")
            else:
                exact[ke] = row[0]

            kb = gloss_key(row[3])
            if kb in base:
                pw, pf = base[kb]
                near.append(f"درس {ln}: «{pw}»=«{pf}» و «{row[0]}»=«{row[3]}» → پایه «{kb}»")
            else:
                base[kb] = (row[0], row[3])

    if same:
        raise SystemExit("معنی فارسیِ یکسان در یک درس:\n  " + "\n  ".join(same))
    if near:
        print("⚠️  معنیِ پایه یکسان (اپ خودش جلوی هم‌زمان آمدنشان در گزینه‌ها را می‌گیرد):")
        for n in near:
            print("     " + n)

    # --- محافظ: هیچ درسی نباید بخش خالی داشته باشد
    #
    # این محافظ نبود و نتیجه‌اش گران تمام شد: درس‌های ۴ تا ۱۰ سطح A2
    # بخش گرامرشان **کاملاً خالی** بود، چون فایل داده‌شان با نامی بود که
    # بارگذار نمی‌شناخت. هیچ خطایی نمی‌داد — نبودِ کلید در دیکشنری خطا
    # نیست — و تنها راه دیدنش باز کردن همان درس روی گوشی بود.
    lessons = [row[0] for row in SYLLABUS]
    holes = []
    for ln in lessons:
        for label, table in (("واژگان", V), ("مکالمه", DLG), ("گرامر", GR), ("تمرین", EX)):
            if not table.get(ln):
                holes.append(f"درس {ln}: بخش «{label}» خالی است")
    if holes:
        raise SystemExit(
            "درسِ ناقص (فایل داده‌اش خوانده نشده یا نوشته نشده):\n  "
            + "\n  ".join(holes)
        )

    wb = Workbook()
    wb.remove(wb.active)

    ws = sheet(wb, "نقشه دروس",
               ["درس", "عنوان انگلیسی", "موضوع گرامر", "ماموریت درس", "تعداد واژه", "رایگان؟", "وضعیت"],
               [7, 30, 30, 28, 11, 9, 10])
    for n, title_en, grammar, theme, free in SYLLABUS:
        ws.append([n, title_en, grammar, theme, len(V.get(n, [])),
                   "بله" if free else "خیر", "آماده" if n in V else "در انتظار محتوا"])
    finish(ws, 7)

    ws = sheet(wb, "واژگان",
               ["شناسه", "درس", "واژه", "بخش کلام", "تلفظ IPA", "معنی فارسی", "جمله مثال", "ترجمه جمله",
                "پرامپت تصویر (کپی کنید)", "نام فایل تصویر", "متن صوتی", "وضعیت"],
               [16, 7, 18, 13, 17, 20, 34, 30, 78, 30, 34, 11])
    for ln in sorted(V):
        for i, (w, pos, ipa, fa, ex, exfa, scene) in enumerate(V[ln], 1):
            slug = w.lower().replace(" ", "_").replace("'", "")
            ws.append([f"{level}-L{ln:02d}-W{i:02d}", ln, w, pos, ipa, fa, ex, exfa,
                       STYLE + scene, f"{low}_l{ln:02d}_w{i:02d}_{slug}.png",
                       f"{w}. {ex}", "آماده"])
    finish(ws, 12, wrap_cols=(7, 8, 9))

    ws = sheet(wb, "مکالمه",
               ["درس", "نوبت", "گوینده", "جمله انگلیسی", "ترجمه فارسی", "صدای TTS", "وضعیت"],
               [7, 8, 11, 42, 38, 11, 10])
    for ln in sorted(DLG):
        for i, (spk, en, fa) in enumerate(DLG[ln], 1):
            ws.append([ln, i, "نفر اول" if spk == "A" else "نفر دوم", en, fa,
                       "زن" if spk == "A" else "مرد", "آماده"])
    finish(ws, 7, wrap_cols=(4, 5))

    ws = sheet(wb, "گرامر",
               ["درس", "کارت", "عنوان", "توضیح فارسی", "مثال‌ها",
                "پرسش بررسی", "پاسخ درست", "گزینه‌ها", "نکته هر گزینه غلط", "وضعیت"],
               [7, 9, 26, 60, 40, 34, 16, 26, 40, 10])
    for ln in sorted(GR):
        for i, card in enumerate(GR[ln], 1):
            c = list(card) + [""] * (7 - len(card))
            ws.append([ln, f"Card {i}", c[0], c[1], c[2], c[3], c[4], c[5], c[6], "آماده"])
    finish(ws, 10, wrap_cols=(4, 5, 6, 8, 9))

    ws = sheet(wb, "تمرین‌ها",
               ["درس", "تمرین", "دسته", "نوع", "صورت سوال", "پاسخ درست", "پاسخ‌های جایگزین",
                "گزینه‌ها", "واژه هدف (لایتنر)", "راهنما", "وضعیت"],
               [7, 11, 11, 17, 40, 26, 26, 32, 18, 32, 10])

    MEANING = {w.lower(): fa for rows in V.values() for (w, _p, _i, fa, *_r) in rows}

    def translate_options(opt, answer, lesson, index):
        out = []
        for w in [x.strip() for x in opt.split("|") if x.strip()]:
            fa = MEANING.get(w.lower())
            if fa is None:
                raise SystemExit(f"درس {lesson} تمرین {index}: معنی «{w}» در جدول واژگان نیست.")
            out.append(fa)
        if answer not in out:
            raise SystemExit(f"درس {lesson} تمرین {index}: پاسخ درست «{answer}» بین گزینه‌ها نیست.")
        return " | ".join(out)

    for ln in sorted(EX):
        for i, (cat, ty, q, ans, alt, opt, tw, h) in enumerate(EX[ln], 1):
            if ty == "تطبیق" and opt:
                opt = translate_options(opt, ans, ln, i)
            ws.append([ln, f"Exercise {i}", cat, ty, q, ans, alt, opt, tw, h, "آماده"])
    finish(ws, 11, wrap_cols=(5, 7, 8, 10))

    if ST:
        ws = sheet(wb, "داستانک",
                   ["درس", "عنوان", "نوع", "متن انگلیسی", "ترجمه فارسی",
                    "پاسخ درست", "گزینه‌ها", "راهنما", "وضعیت"],
                   [7, 26, 9, 60, 46, 22, 34, 30, 10])
        for ln in sorted(ST):
            title, paras, questions = ST[ln]
            for en, fa in paras:
                ws.append([ln, title, "بند", en, fa, "", "", "", "آماده"])
            for q_en, q_fa, ans, opts, hint in questions:
                ws.append([ln, title, "پرسش", q_en, q_fa, ans, " | ".join(opts), hint, "آماده"])
        finish(ws, 9, wrap_cols=(4, 5, 7, 8))

    wb.save(out)

    nv = sum(len(x) for x in V.values())
    print(f"ساخته شد: {out}")
    print(f"  سرفصل  : {len(SYLLABUS)} درس")
    print(f"  واژگان : {nv} ردیف  ({len(V)} درس نوشته‌شده)  ← {nv} تصویر لازم است")
    print(f"  مکالمه : {sum(len(x) for x in DLG.values())} ردیف")
    print(f"  گرامر  : {sum(len(x) for x in GR.values())} کارت")
    print(f"  تمرین  : {sum(len(x) for x in EX.values())} تمرین")
    if ST:
        nq = sum(len(v[2]) for v in ST.values())
        print(f"  داستانک: {len(ST)} داستان · {nq} پرسش درک مطلب")
    if EX:
        tp = sum(1 for rows in EX.values() for e in rows if e[0] == P)
        tt = sum(len(rows) for rows in EX.values())
        print(f"\n  بررسی قاعده ۶۰٪ تولیدی: {tp}/{tt} = {tp*100//tt}%", "✅" if tp / tt >= 0.60 else "❌")
    missing = [n for n, *_ in SYLLABUS if n not in V]
    if missing:
        print(f"\n  ⏳ دروس بدون محتوا: {', '.join(map(str, missing))}")


if __name__ == "__main__":
    main()
